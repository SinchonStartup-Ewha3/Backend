from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path

from django.core.management.base import BaseCommand, CommandError
from django.db import transaction
from openpyxl import load_workbook

from core.models import Region

DEFAULT_PATH = Path(__file__).resolve().parents[2] / "data" / "kma_grid.xlsx"
SIX_PLACES = Decimal("0.000001")

REQUIRED_COLUMNS = [
    "행정구역코드", "1단계", "2단계", "3단계",
    "격자 X", "격자 Y", "경도(초/100)", "위도(초/100)",
]


def to_decimal(value):
    # 모델이 소수점 6자리라 반올림해서 맞춰준다
    return Decimal(str(value)).quantize(SIX_PLACES, rounding=ROUND_HALF_UP)


class Command(BaseCommand):
    help = "기상청 격자 엑셀 파일로 Region 데이터를 적재합니다."

    def add_arguments(self, parser):
        parser.add_argument("--path", default=str(DEFAULT_PATH))

    def handle(self, *args, **options):
        path = Path(options["path"])
        if not path.exists():
            raise CommandError(f"파일을 찾을 수 없습니다: {path}")

        wb = load_workbook(path, read_only=True, data_only=True)
        rows = wb.worksheets[0].iter_rows(values_only=True)

        header = [str(h).strip() if h is not None else "" for h in next(rows)]
        col = {name: i for i, name in enumerate(header)}

        missing = [c for c in REQUIRED_COLUMNS if c not in col]
        if missing:
            raise CommandError(f"엑셀에서 찾을 수 없는 컬럼: {missing}\n실제 헤더: {header}")

        regions = []
        for row in rows:
            code = row[col["행정구역코드"]]
            lat = row[col["위도(초/100)"]]
            lng = row[col["경도(초/100)"]]
            if code is None or lat is None or lng is None:
                continue

            if isinstance(code, float):  # 1100000000.0 같은 형태 방지
                code = int(code)

            name = " ".join(
                str(row[col[level]]).strip()
                for level in ("1단계", "2단계", "3단계")
                if row[col[level]]
            )

            regions.append(Region(
                region_code=str(code).strip(),
                region_name=name,
                lat=to_decimal(lat),
                lng=to_decimal(lng),
                grid_nx=int(row[col["격자 X"]]),
                grid_ny=int(row[col["격자 Y"]]),
            ))

        with transaction.atomic():
            Region.objects.bulk_create(
                regions,
                update_conflicts=True,
                unique_fields=["region_code"],
                update_fields=["region_name", "lat", "lng", "grid_nx", "grid_ny"],
            )

        self.stdout.write(self.style.SUCCESS(f"Region {len(regions)}건 적재 완료"))